from copy import deepcopy
from io import BytesIO

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from pypdf import PdfReader, PdfWriter

from backend.resumes.canvas_schemas import CanvasStyle
from backend.resumes.content import build_content
from backend.resumes.quality import (
    ResumeQualityError,
    extract_docx_text_in_order,
    validate_resume_artifacts,
)
from backend.resumes.renderers import fonts
from backend.resumes.renderers.ats import render_ats_docx, render_ats_pdf
from backend.resumes.renderers.links import is_safe_url
from backend.resumes.templates import PRESETS


def test_catalog_defaults_are_deeply_immutable_and_valid():
    for preset in PRESETS:
        CanvasStyle.model_validate(dict(preset.preview_style))
        with pytest.raises(TypeError):
            preset.preview_style["accent_color"] = "#FFFFFF"
        with pytest.raises(TypeError):
            preset.section_headings["experience"] = "Changed"


def test_initial_presentation_preserves_explicit_canvas_style():
    from backend.resumes.canvas_schemas import CanvasBlock, CanvasSection, ResumeCanvasDocument
    from backend.resumes.presentation import apply_presentation
    from backend.resumes.templates import get_template_preset

    canvas = ResumeCanvasDocument(
        sections=[
            CanvasSection(
                id="identity",
                kind="identity",
                title="IDENTITY",
                blocks=[CanvasBlock(id="identity-main", kind="identity")],
            )
        ],
        style=CanvasStyle(base_font_size=11, accent_color="#123456", margin_mm=22),
    )
    changed = apply_presentation(canvas, get_template_preset("swiss-software-de"), locale="de")
    assert changed.style.base_font_size == 11
    assert changed.style.accent_color == "#123456"
    assert changed.style.margin_mm == 22
    assert changed.style.columns == 2
    assert changed.sections[0].blocks == canvas.sections[0].blocks


@pytest.mark.parametrize(
    "url",
    [
        "https:relative",
        "https://user:secret@example.test",
        "https://@example.test",
        " https://example.test",
        "https://example.test/a b",
        "https://[broken",
        "https://example.test:99999",
        "mailto:test@example.test?body=private",
        "https://example.test/%0A",
    ],
)
def test_unsafe_links_are_plain_text(url):
    assert not is_safe_url(url)


def test_unicode_font_unavailability_fails_clearly(monkeypatch):
    monkeypatch.setattr(fonts, "_FONTS_REGISTERED", False)
    monkeypatch.setattr(fonts, "_candidate_font_paths", lambda: [])
    with pytest.raises(ResumeQualityError, match="local Unicode font is unavailable"):
        fonts.ensure_unicode_font_registered()


def test_docx_traversal_retains_all_nested_and_merged_cells():
    document = Document()
    table = document.add_table(rows=40, cols=2)
    expected = []
    for index, row in enumerate(table.rows):
        for column, cell in enumerate(row.cells):
            text = f"Marker-{index}-{column}"
            cell.text = text
            expected.append(text)
    nested = table.cell(0, 0).add_table(rows=1, cols=2)
    nested.cell(0, 0).merge(nested.cell(0, 1)).text = "Nested unique"
    text = extract_docx_text_in_order(document)
    assert all(text.count(marker + "\n") == 1 for marker in expected[:-1])
    assert text.count("Nested unique") == 1
    assert text.endswith(expected[-1])


@pytest.mark.parametrize("preset", PRESETS, ids=lambda preset: preset.id)
def test_every_preset_publishes_complete_unicode_content_and_links(
    preset,
    client,
    auth_headers,
    detailed_profile_payload,
    db_session,
):
    payload = deepcopy(detailed_profile_payload)
    payload["display_name"] = "Łucja Müller"
    payload["location"] = {"city": "Łódź", "country": "PL"}
    payload["summary"] = "Entwickelt zuverlässige Systeme in Zürich."
    payload["facts"] = payload["facts"][:1]
    profile = client.put("/api/v1/career-profile", json=payload, headers=auth_headers).json()
    created = client.post(
        "/api/v1/resumes",
        json={
            "title": "Synthetic publication",
            "template_id": preset.id,
            "template_version": 1,
            "selected_fact_ids": [profile["facts"][0]["id"]],
        },
        headers=auth_headers,
    )
    assert created.status_code == 201, created.text
    draft = created.json()
    assert draft["template_layout"] == preset.layout
    response = client.post(f"/api/v1/resumes/{draft['id']}/publish", json={}, headers=auth_headers)
    assert response.status_code == 201, response.text
    version = response.json()
    assert version["quality_report"]["max_pages"] == preset.page_budget
    from backend.resumes.models import ResumeVersion

    stored = db_session.get(ResumeVersion, version["id"])
    assert stored.snapshot["resume"]["draft_revision"] == draft["revision"]
    assert preset.section_headings["experience"] in version["quality_report"]["required_headings"]
    artifacts = version["artifacts"]
    downloaded = {}
    for artifact in artifacts:
        result = client.get(f"/api/v1/resume-artifacts/{artifact['id']}", headers=auth_headers)
        assert result.status_code == 200, result.text
        downloaded[artifact["format"]] = result.content
    pdf = PdfReader(BytesIO(downloaded["pdf"]))
    text = "\n".join(page.extract_text() for page in pdf.pages)
    for value in [
        "Łucja Müller",
        "Łódź",
        "Zürich",
        payload["summary"],
        payload["facts"][0]["payload"]["description"],
        "Reduced deployment lead time by 40%.",
    ]:
        assert value in text
    urls = [
        str(annotation.get_object().get("/A", {}).get("/URI", ""))
        for page in pdf.pages
        for annotation in page.get("/Annots", [])
    ]
    assert payload["website"] in urls
    word = Document(BytesIO(downloaded["docx"]))
    assert (
        str(word.styles["Heading 2"].font.color.rgb)
        == preset.preview_style["accent_color"].lstrip("#").upper()
    )
    assert payload["website"] in [
        relation.target_ref
        for relation in word.part.rels.values()
        if relation.reltype == RELATIONSHIP_TYPE.HYPERLINK
    ]
    assert "Łucja Müller" in extract_docx_text_in_order(word)
    assert all(
        abs(section.page_width.mm - 210) < 0.2 and abs(section.page_height.mm - 297) < 0.2
        for section in word.sections
    )


def test_pdf_non_a4_is_rejected():
    snapshot = {
        "profile": {"display_name": "Synthetic", "headline": "", "summary": ""},
        "resume": {"section_config": {"order": ["skill"]}},
        "facts": [{"id": "fact", "fact_type": "skill", "payload": {"name": "Python"}}],
    }
    original = PdfReader(BytesIO(render_ats_pdf(snapshot)))
    original.pages[0].mediabox.upper_right = (612, 792)
    writer = PdfWriter()
    writer.add_page(original.pages[0])
    output = BytesIO()
    writer.write(output)
    content = build_content(snapshot)
    with pytest.raises(ResumeQualityError, match="A4"):
        validate_resume_artifacts(
            pdf=output.getvalue(),
            docx=render_ats_docx(snapshot),
            required_headings=content.required_headings,
            required_text=["Synthetic", "Python"],
            template_kind="ats",
            expect_photo=False,
        )


def _write_payload(draft, **changes):
    result = {
        key: deepcopy(draft[key])
        for key in (
            "title",
            "template_id",
            "template_version",
            "locale",
            "template_kind",
            "section_config",
            "selected_fact_ids",
            "content_overrides",
            "canvas_document",
            "photo_asset_id",
        )
    }
    result["expected_revision"] = draft["revision"]
    result.update(changes)
    return result


def test_preset_switch_retains_manual_multi_fact_identity_and_selection_changes(
    client,
    auth_headers,
    saved_detailed_profile,
):
    facts = saved_detailed_profile["facts"]
    created = client.post(
        "/api/v1/resumes",
        json={"title": "Evidence", "selected_fact_ids": [f["id"] for f in facts]},
        headers=auth_headers,
    )
    draft = created.json()
    canvas = draft["canvas_document"]
    experience = next(s for s in canvas["sections"] if s["kind"] == "experience")
    education = next(s for s in canvas["sections"] if s["kind"] == "education")
    first = experience["blocks"][0]
    second_id = education["blocks"][0]["fact_ids"][0]
    first["fact_ids"].append(second_id)
    first["content"]["title"] = "Approved multidisciplinary claim"
    first["manual_fields"] = ["title"]
    experience["title"] = "Selected systems"
    canvas["sections"].remove(education)
    saved = client.put(
        f"/api/v1/resumes/{draft['id']}", json=_write_payload(draft), headers=auth_headers
    )
    assert saved.status_code == 200, saved.text
    original = saved.json()
    switched = client.put(
        f"/api/v1/resumes/{draft['id']}",
        json=_write_payload(original, template_id="swiss-software-de", locale="de"),
        headers=auth_headers,
    )
    assert switched.status_code == 200, switched.text
    result = switched.json()
    assert [
        block for section in result["canvas_document"]["sections"] for block in section["blocks"]
    ] == [
        block for section in original["canvas_document"]["sections"] for block in section["blocks"]
    ]
    assert (
        next(s for s in result["canvas_document"]["sections"] if s["kind"] == "experience")["title"]
        == "Selected systems"
    )
    selected = [fid for fid in result["selected_fact_ids"] if fid != second_id]
    overrides = {
        fid: value for fid, value in result["content_overrides"].items() if fid in selected
    }
    removed = client.put(
        f"/api/v1/resumes/{draft['id']}",
        json=_write_payload(result, selected_fact_ids=selected, content_overrides=overrides),
        headers=auth_headers,
    )
    assert removed.status_code == 200, removed.text
    blocks = [
        block
        for section in removed.json()["canvas_document"]["sections"]
        for block in section["blocks"]
    ]
    assert all(second_id not in block["fact_ids"] for block in blocks)
    assert (
        next(block for block in blocks if block["id"] == first["id"])["content"]["title"]
        == first["content"]["title"]
    )


def test_ats_suppresses_but_preserves_photo_for_roundtrip(
    client, auth_headers, saved_detailed_profile
):
    from PIL import Image

    photo = BytesIO()
    Image.new("RGB", (200, 200), "navy").save(photo, format="PNG")
    uploaded = client.post(
        "/api/v1/career-profile/photo",
        files={"file": ("portrait.png", photo.getvalue(), "image/png")},
        headers=auth_headers,
    )
    assert uploaded.status_code == 201
    photo_id = uploaded.json()["id"]
    created = client.post(
        "/api/v1/resumes",
        json={
            "title": "Photo roundtrip",
            "template_id": "swiss-software-en",
            "selected_fact_ids": [saved_detailed_profile["facts"][0]["id"]],
            "photo_asset_id": photo_id,
        },
        headers=auth_headers,
    )
    assert created.status_code == 201, created.text
    draft = created.json()
    canvas = deepcopy(draft["canvas_document"])
    canvas["style"]["columns"] = 1
    switched = client.put(
        f"/api/v1/resumes/{draft['id']}",
        json=_write_payload(
            draft, template_id="software-en", photo_asset_id=None, canvas_document=canvas
        ),
        headers=auth_headers,
    )
    assert switched.status_code == 200, switched.text
    assert switched.json()["photo_asset_id"] == photo_id
    published = client.post(f"/api/v1/resumes/{draft['id']}/publish", headers=auth_headers)
    assert published.status_code == 201, published.text
    assert published.json()["quality_report"]["pdf_image_count"] == 0
    returned = client.put(
        f"/api/v1/resumes/{draft['id']}",
        json=_write_payload(switched.json(), template_id="swiss-software-en"),
        headers=auth_headers,
    )
    assert returned.status_code == 200, returned.text
    assert returned.json()["photo_asset_id"] == photo_id


def test_external_generation_metadata_roundtrips_and_bounds():
    from uuid import uuid4

    from pydantic import ValidationError

    from backend.resumes.canvas_schemas import GenerationContext

    legacy = GenerationContext(source_profile_revision=1)
    assert legacy.mode == "deterministic"
    value = GenerationContext(
        mode="external-agent",
        source_profile_revision=2,
        request_id=str(uuid4()),
        grant_id=str(uuid4()),
        input_digest="1" * 64,
        payload_digest="2" * 64,
        claim_evidence_map={"block": [str(uuid4())]},
    )
    assert GenerationContext.model_validate(value.model_dump()) == value
    with pytest.raises(ValidationError):
        GenerationContext(source_profile_revision=1, request_id="unbounded" * 100)
    with pytest.raises(ValidationError):
        GenerationContext(source_profile_revision=1, payload_digest="not-a-digest")


def test_publication_and_restore_preserve_external_provenance(
    client,
    auth_headers,
    saved_detailed_profile,
    db_session,
):
    from uuid import uuid4

    from backend.resumes.models import ResumeDraft, ResumeVersion

    fact_id = saved_detailed_profile["facts"][0]["id"]
    draft = client.post(
        "/api/v1/resumes",
        json={"title": "Provenance", "selected_fact_ids": [fact_id]},
        headers=auth_headers,
    ).json()
    metadata = {
        "mode": "external-agent",
        "source_profile_revision": saved_detailed_profile["revision"],
        "request_id": str(uuid4()),
        "grant_id": str(uuid4()),
        "input_digest": "a" * 64,
        "payload_digest": "b" * 64,
        "claim_evidence_map": {"claim": [fact_id]},
    }
    stored = db_session.get(ResumeDraft, draft["id"])
    stored.generation_context = deepcopy(metadata)
    db_session.commit()
    publication = client.post(f"/api/v1/resumes/{draft['id']}/publish", headers=auth_headers)
    assert publication.status_code == 201, publication.text
    version = db_session.get(ResumeVersion, publication.json()["id"])
    assert version.snapshot["resume"]["generation_context"] == metadata
    assert version.snapshot["resume"]["draft_revision"] == draft["revision"]
    restored = client.post(
        f"/api/v1/resumes/{draft['id']}/versions/{version.id}/restore",
        json={"expected_revision": draft["revision"]},
        headers=auth_headers,
    )
    assert restored.status_code == 200, restored.text
    assert {key: restored.json()["generation_context"][key] for key in metadata} == metadata


@pytest.mark.parametrize("preset_id", ["software-en", "swiss-software-de", "operational-de"])
@pytest.mark.parametrize("hidden", ["section", "block", "cleared"])
def test_export_does_not_restore_hidden_identity_from_profile(
    client,
    auth_headers,
    saved_detailed_profile,
    preset_id,
    hidden,
):
    created = client.post(
        "/api/v1/resumes",
        json={
            "title": "Hidden identity",
            "template_id": preset_id,
            "selected_fact_ids": [saved_detailed_profile["facts"][0]["id"]],
        },
        headers=auth_headers,
    )
    assert created.status_code == 201, created.text
    draft = created.json()
    identity = next(
        section for section in draft["canvas_document"]["sections"] if section["kind"] == "identity"
    )
    block = identity["blocks"][0]
    hidden_values = [block["content"]["subtitle"], saved_detailed_profile["email"]]
    if hidden == "section":
        identity["visible"] = False
    elif hidden == "block":
        block["visible"] = False
    else:
        block["content"].update(subtitle="", description="")
        block["manual_fields"] = ["subtitle", "description"]
    saved = client.put(
        f"/api/v1/resumes/{draft['id']}", json=_write_payload(draft), headers=auth_headers
    )
    assert saved.status_code == 200, saved.text
    content = build_content(
        {"profile": saved_detailed_profile, "resume": saved.json(), "facts": []}
    )
    assert content.headline == "" and content.contact_line == ""
    published = client.post(f"/api/v1/resumes/{draft['id']}/publish", headers=auth_headers)
    if hidden in {"section", "block"}:
        assert content.display_name == ""
        assert published.status_code == 422
        assert published.json()["detail"] == "A publishable resume requires a visible identity name"
        return
    assert published.status_code == 201, published.text
    for artifact in published.json()["artifacts"]:
        content = client.get(
            f"/api/v1/resume-artifacts/{artifact['id']}", headers=auth_headers
        ).content
        if artifact["format"] == "pdf":
            reader = PdfReader(BytesIO(content))
            text = "\n".join(page.extract_text() for page in reader.pages)
            assert reader.metadata.title == saved_detailed_profile["display_name"]
        else:
            document = Document(BytesIO(content))
            text = extract_docx_text_in_order(document)
            assert document.core_properties.title == saved_detailed_profile["display_name"]
        assert "Principal Engineer" in text
        assert all(value not in text for value in hidden_values if value)
