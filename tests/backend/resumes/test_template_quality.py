import io
import zipfile
from xml.etree import ElementTree

import pytest
from docx import Document
from pypdf import PdfReader

from backend.resumes.quality import (
    ResumeQualityError,
    check_no_placeholders,
    extract_docx_text_in_order,
    validate_resume_artifacts,
)
from backend.resumes.renderers.ats import render_ats_docx, render_ats_pdf
from backend.resumes.renderers.operational import render_operational_docx, render_operational_pdf
from backend.resumes.renderers.photo import render_photo_docx, render_photo_pdf
from backend.resumes.templates import PRESETS, get_template_preset


def _sample_snapshot(preset_id: str, extra_facts=None, text_override=None) -> dict:
    preset = get_template_preset(preset_id)
    profile = {
        "display_name": text_override.get("name", "Mira Vale") if text_override else "Mira Vale",
        "headline": "Lead Systems Architect",
        "summary": "Specialist in reliable distributed systems and local computation.",
        "email": "mira.vale@example.ch",
        "phone": "+41 44 123 45 67",
        "location": {"city": "Zürich", "country": "CH"},
        "website": "https://miravale.ch",
        "linkedin": "https://linkedin.com/in/miravale",
        "github": "https://github.com/miravale",
    }
    facts = extra_facts or [
        {
            "id": "fact-1",
            "fact_type": "experience",
            "position": 0,
            "payload": {
                "role": "Systems Architect",
                "organization": "Alpine Cloud Systems",
                "start_date": "2021-03",
                "current": True,
                "description": "Architected low-latency message streaming pipelines.",
                "achievements": ["Reduced latency by 45%."],
            },
            "verification_status": "confirmed",
            "source_document_id": None,
        },
        {
            "id": "fact-2",
            "fact_type": "education",
            "position": 1,
            "payload": {
                "qualification": "MSc Computer Science",
                "institution": "ETH Zürich",
                "start_date": "2018-09",
                "end_date": "2020-09",
                "description": "Distributed systems focus.",
            },
            "verification_status": "confirmed",
            "source_document_id": None,
        },
        {
            "id": "fact-3",
            "fact_type": "skill",
            "position": 2,
            "payload": {"name": "Python", "category": "Backend"},
            "verification_status": "confirmed",
            "source_document_id": None,
        },
    ]
    return {
        "profile": profile,
        "resume": {
            "title": f"Test Resume ({preset_id})",
            "template_kind": preset.template_kind,
            "template_id": preset.id,
            "template_version": preset.version,
            "locale": preset.locale,
            "section_config": {
                "order": ["experience", "education", "skill"],
                "include_summary": True,
                "include_email": True,
                "include_phone": True,
                "include_location": True,
                "include_links": True,
            },
            "content_overrides": {},
            "canvas_document": {
                "style": {"columns": preset.preview_style.get("columns", 1)},
                "sections": [],
            },
            "generation_context": {},
        },
        "facts": facts,
        "selected_fact_ids": [f["id"] for f in facts],
    }


def test_renderers_generate_valid_a4_docx_and_pdf_for_all_nine_presets():
    for preset in PRESETS:
        snapshot = _sample_snapshot(preset.id)
        if preset.layout == "ats":
            pdf = render_ats_pdf(snapshot)
            docx = render_ats_docx(snapshot)
        elif preset.layout == "swiss-operational":
            pdf = render_operational_pdf(snapshot, None)
            docx = render_operational_docx(snapshot, None)
        else:
            pdf = render_photo_pdf(snapshot, None)
            docx = render_photo_docx(snapshot, None)

        assert len(pdf) > 1000
        assert len(docx) > 1000

        # Verify PDF page dimensions (A4 = 595.27 x 841.89 points)
        pdf_reader = PdfReader(io.BytesIO(pdf))
        for page in pdf_reader.pages:
            w = float(page.mediabox.width)
            h = float(page.mediabox.height)
            assert abs(w - 595.27) < 3.0
            assert abs(h - 841.89) < 3.0
        assert len(pdf_reader.pages) <= preset.page_budget

        # Verify DOCX section page dimensions (Mm(210) x Mm(297))
        docx_doc = Document(io.BytesIO(docx))
        for sec in docx_doc.sections:
            assert abs(sec.page_width.mm - 210.0) < 1.0
            assert abs(sec.page_height.mm - 297.0) < 1.0

        # Validate with quality engine
        from backend.resumes.content import build_content

        content = build_content(snapshot)
        required_text = [
            content.display_name,
            content.headline,
            *content.contact_line.split(" | "),
            content.summary,
        ]
        for section in content.sections:
            for entry in section.entries:
                required_text.extend(
                    [
                        entry.title,
                        entry.subtitle,
                        entry.date_range,
                        entry.description,
                        *entry.bullets,
                    ]
                )
        required_text = [text for text in required_text if text]

        report = validate_resume_artifacts(
            pdf=pdf,
            docx=docx,
            required_headings=content.required_headings,
            required_text=required_text,
            template_kind=preset.template_kind,
            expect_photo=False,
            columns=preset.preview_style.get("columns", 1),
            page_budget=preset.page_budget,
        )
        assert report["passed"] is True
        assert report["within_page_limit"] is True


def test_unicode_font_preserves_polish_and_german_diacritics_in_pdf():
    snapshot = _sample_snapshot(
        "software-en",
        text_override={"name": "Stanisław Łódź-Kraków"},
        extra_facts=[
            {
                "id": "fact-pl",
                "fact_type": "experience",
                "position": 0,
                "payload": {
                    "role": "Lead Engineer",
                    "organization": "Polskie Sieci Łódź",
                    "start_date": "2020-01",
                    "current": True,
                    "description": "Praca w Krakowie przy Straßenbahn Zürich.",
                    "achievements": ["Zwiększono wydajność o 30%."],
                },
                "verification_status": "confirmed",
                "source_document_id": None,
            }
        ],
    )
    pdf = render_ats_pdf(snapshot)
    reader = PdfReader(io.BytesIO(pdf))
    extracted = "\n".join(p.extract_text() or "" for p in reader.pages)

    for value in (
        "Stanisław Łódź-Kraków",
        "Straßenbahn Zürich",
        "Polskie Sieci Łódź",
        "Zwiększono wydajność o 30%.",
    ):
        assert value in extracted


def test_docx_text_extraction_traverses_tables_in_document_order():
    snapshot = _sample_snapshot("swiss-software-en")
    docx = render_photo_docx(snapshot, None)
    doc = Document(io.BytesIO(docx))
    text = extract_docx_text_in_order(doc)

    assert "Mira Vale" in text
    assert "Alpine Cloud Systems" in text
    assert "ETH Zürich" in text


def test_real_hyperlinks_in_pdf_and_docx():
    snapshot = _sample_snapshot("software-en")
    pdf = render_ats_pdf(snapshot)
    docx = render_ats_docx(snapshot)

    # DOCX: check XML for w:hyperlink
    with zipfile.ZipFile(io.BytesIO(docx)) as archive:
        doc_xml = archive.read("word/document.xml").decode("utf-8")
        relationships = ElementTree.fromstring(archive.read("word/_rels/document.xml.rels"))
    assert "w:hyperlink" in doc_xml
    relationship_targets = {relationship.get("Target") for relationship in relationships}
    assert {
        snapshot["profile"]["website"],
        snapshot["profile"]["github"],
    } & relationship_targets

    # PDF: check PDF annotations or links
    reader = PdfReader(io.BytesIO(pdf))
    urls = [
        str(annotation.get_object().get("/A", {}).get("/URI", ""))
        for page in reader.pages
        for annotation in page.get("/Annots", [])
    ]
    assert snapshot["profile"]["website"] in urls
    assert snapshot["profile"]["github"] in urls


def test_placeholder_detection_blocks_unresolved_placeholders():
    bad_texts = [
        "Worked at [COMPANY] in 2022",
        "Role: [ROLE] at Acme",
        "[TODO] add more details",
        "TODO: replace this with achievements",
        "FIXME: verify numbers",
    ]
    for bad in bad_texts:
        with pytest.raises(ResumeQualityError, match="Unresolved placeholder detected"):
            check_no_placeholders([bad])

    clean_texts = ["Worked at Acme Corp in 2022", "Role: Engineer at Google"]
    check_no_placeholders(clean_texts)  # Should not raise


def test_operational_preset_enforces_single_page_budget():
    preset = get_template_preset("operational-de")
    assert preset.page_budget == 1

    snapshot = _sample_snapshot("operational-de")
    pdf = render_operational_pdf(snapshot, None)
    docx = render_operational_docx(snapshot, None)

    reader = PdfReader(io.BytesIO(pdf))
    assert len(reader.pages) == 1

    from backend.resumes.content import build_content

    content = build_content(snapshot)
    required_text = [content.display_name]
    for section in content.sections:
        for entry in section.entries:
            if entry.title:
                required_text.append(entry.title)

    report = validate_resume_artifacts(
        pdf=pdf,
        docx=docx,
        required_headings=content.required_headings,
        required_text=required_text,
        template_kind="photo",
        expect_photo=False,
        columns=2,
        page_budget=1,
    )
    assert report["passed"] is True
    assert report["within_page_limit"] is True


@pytest.mark.parametrize("preset_id", ["software-en", "swiss-software-de", "operational-de"])
def test_long_content_is_preserved_and_page_overflow_rejected(preset_id):
    snapshot = _sample_snapshot(preset_id)
    snapshot["facts"] = [
        {
            **snapshot["facts"][0],
            "id": f"fact-{index}",
            "position": index,
            "payload": {
                "role": f"Role marker {index:02d}",
                "organization": "Synthetic Systems",
                "description": "Documented dependable local systems. " * 40,
                "achievements": [f"Final evidence marker {index:02d}."],
            },
        }
        for index in range(12)
    ]
    preset = get_template_preset(preset_id)
    pdf = render_ats_pdf(snapshot) if preset.layout == "ats" else render_photo_pdf(snapshot, None)
    docx = (
        render_ats_docx(snapshot) if preset.layout == "ats" else render_photo_docx(snapshot, None)
    )
    reader = PdfReader(io.BytesIO(pdf))
    extracted = "\n".join(page.extract_text() or "" for page in reader.pages)
    word_text = extract_docx_text_in_order(Document(io.BytesIO(docx)))
    for index in range(12):
        for text in (extracted, word_text):
            assert f"Role marker {index:02d}" in text
            assert f"Final evidence marker {index:02d}." in text
    assert len(reader.pages) > preset.page_budget
    with pytest.raises(ResumeQualityError, match="configured limit"):
        validate_resume_artifacts(
            pdf=pdf,
            docx=docx,
            required_headings=[],
            required_text=[],
            template_kind=preset.template_kind,
            expect_photo=False,
            page_budget=preset.page_budget,
        )
