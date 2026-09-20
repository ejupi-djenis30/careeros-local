from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory

import pytest
from docx import Document
from pydantic import ValidationError

from backend.career.reference_parsing import (
    extract_docx_text_in_document_order,
    parse_goal_preferences,
    validate_safe_source_input,
)
from backend.career.schemas import PreferenceCandidate
from backend.career.source_parsing import SourceImportError


def _profile_payload(expected_revision=0):
    return {
        "expected_revision": expected_revision,
        "display_name": "Fictional Candidate",
        "headline": "Software Architect",
        "summary": "Builds robust systems.",
        "email": "candidate@example.test",
        "phone": "+41 79 111 22 33",
        "location": {"city": "Zurich", "country": "CH"},
        "preferences": {},
        "facts": [],
        "goals": [],
    }


def test_fictional_profile_narrative_and_goal_import_paths(client, auth_headers, monkeypatch):
    """Test full fictional Profile.md, Storytelling.md, and Goal.md import lifecycle."""
    profile_res = client.put("/api/v1/career-profile", json=_profile_payload(0), headers=auth_headers)
    assert profile_res.status_code == 200, profile_res.text

    with TemporaryDirectory() as temp_dir:
        monkeypatch.setattr("backend.career.sources.settings.DATA_DIR", temp_dir)

        # 1. Profile.md import (default role: profile)
        profile_md = b"# Fictional Profile\n\nCompetenze: Python, FastAPI, TypeScript\n\nDesigned distributed consensus protocol."
        res_profile = client.post(
            "/api/v1/career-profile/sources",
            headers=auth_headers,
            data={"source_role": "profile"},
            files={"file": ("Profile.md", profile_md, "text/markdown")},
        )
        assert res_profile.status_code == 201, res_profile.text
        data_p = res_profile.json()
        assert data_p["source_role"] == "profile"
        assert len(data_p["candidates"]) > 0
        assert len(data_p["preference_candidates"]) == 0
        assert any(c["payload"].get("name") == "Python" for c in data_p["candidates"])

        # 2. Storytelling.md import (role: narrative)
        story_md = (
            b"# Career Narrative\n\n"
            b"I started programming on early microcomputers and discovered a passion for resilient architectures.\n\n"
            b"Over the years, I mentored dozens of engineers and advocated for clean domain models."
        )
        res_story = client.post(
            "/api/v1/career-profile/sources",
            headers=auth_headers,
            data={"source_role": "narrative"},
            files={"file": ("Storytelling.md", story_md, "text/markdown")},
        )
        assert res_story.status_code == 201, res_story.text
        data_s = res_story.json()
        assert data_s["source_role"] == "narrative"
        assert len(data_s["candidates"]) == 0
        assert len(data_s["preference_candidates"]) == 0
        assert any("narrative text does not generate career facts" in note for note in data_s["review_notes"])

        # 3. Goal.md import (role: goals)
        goal_md = (
            b"# Career Goals & Preferences\n\n"
            b"Target roles: Principal Engineer, Solutions Architect\n"
            b"Preferred locations: Zurich, Bern\n"
            b"Preferred languages: English, German\n"
            b"Preferred work modes: hybrid, remote\n"
            b"Contract types: permanent, contract\n"
            b"Workload bounds: 80 - 100%\n"
            b"Remote only: false\n"
            b"Maximum commute distance: 45 km\n"
            b"Available from: 2026-11-01\n"
            b"Notice period days: 60\n\n"
            b"My ambition is to lead architectural governance across international teams."
        )
        res_goal = client.post(
            "/api/v1/career-profile/sources",
            headers=auth_headers,
            data={"source_role": "goals"},
            files={"file": ("Goal.md", goal_md, "text/markdown")},
        )
        assert res_goal.status_code == 201, res_goal.text
        data_g = res_goal.json()
        assert data_g["source_role"] == "goals"
        assert len(data_g["candidates"]) == 0, "Goals must never produce career fact candidates or achievements"
        assert len(data_g["preference_candidates"]) >= 8

        pref_fields = {pc["field"]: pc["value"] for pc in data_g["preference_candidates"]}
        assert pref_fields["target_roles"] == ["Principal Engineer", "Solutions Architect"]
        assert pref_fields["preferred_locations"] == ["Zurich", "Bern"]
        assert pref_fields["preferred_languages"] == ["English", "German"]
        assert pref_fields["preferred_work_modes"] == ["hybrid", "remote"]
        assert pref_fields["contract_types"] == ["permanent", "contract"]
        assert pref_fields["workload_min"] == 80
        assert pref_fields["workload_max"] == 100
        assert pref_fields["remote_only"] is False
        assert pref_fields["hard_max_distance_km"] == 45.0
        assert pref_fields["available_from"] == "2026-11-01"
        assert pref_fields["notice_period_days"] == 60

        # Verify prose explanation in review notes
        assert any("goals do not generate career achievements" in note for note in data_g["review_notes"])


def test_template_reference_role_produces_no_facts(client, auth_headers, monkeypatch):
    """Test that template_reference role preserves layout guidance without creating facts."""
    profile_res = client.put("/api/v1/career-profile", json=_profile_payload(0), headers=auth_headers)
    assert profile_res.status_code == 200

    with TemporaryDirectory() as temp_dir:
        monkeypatch.setattr("backend.career.sources.settings.DATA_DIR", temp_dir)
        content = b"Layout guidance: place summary at top, two-column skills grid, compact bullet lists."
        res = client.post(
            "/api/v1/career-profile/sources",
            headers=auth_headers,
            data={"source_role": "template_reference"},
            files={"file": ("Template.txt", content, "text/plain")},
        )
        assert res.status_code == 201
        data = res.json()
        assert data["source_role"] == "template_reference"
        assert len(data["candidates"]) == 0
        assert len(data["preference_candidates"]) == 0
        assert any("layout guidance" in note for note in data["review_notes"])


def test_invalid_source_role_rejected(client, auth_headers):
    """Test that unknown or invalid source_role is rejected with 422."""
    client.put("/api/v1/career-profile", json=_profile_payload(0), headers=auth_headers)
    res = client.post(
        "/api/v1/career-profile/sources",
        headers=auth_headers,
        data={"source_role": "unsupported_role"},
        files={"file": ("file.txt", b"Hello", "text/plain")},
    )
    assert res.status_code == 422
    assert "Invalid source_role" in res.text


def test_same_bytes_same_role_reuses_source_conflicting_role_fails(client, auth_headers, monkeypatch):
    """Test same bytes with same role reuses source; conflicting role returns explicit 409 error."""
    client.put("/api/v1/career-profile", json=_profile_payload(0), headers=auth_headers)
    with TemporaryDirectory() as temp_dir:
        monkeypatch.setattr("backend.career.sources.settings.DATA_DIR", temp_dir)
        content = b"Shared document content for reuse testing."

        first = client.post(
            "/api/v1/career-profile/sources",
            headers=auth_headers,
            data={"source_role": "profile"},
            files={"file": ("doc.txt", content, "text/plain")},
        )
        assert first.status_code == 201
        first_data = first.json()

        # Same bytes, same role -> reuses existing record
        second = client.post(
            "/api/v1/career-profile/sources",
            headers=auth_headers,
            data={"source_role": "profile"},
            files={"file": ("doc_copy.txt", content, "text/plain")},
        )
        assert second.status_code == 201
        assert second.json()["id"] == first_data["id"]

        # Same bytes, conflicting role -> raises explicit 409 conflict
        conflicting = client.post(
            "/api/v1/career-profile/sources",
            headers=auth_headers,
            data={"source_role": "narrative"},
            files={"file": ("doc.txt", content, "text/plain")},
        )
        assert conflicting.status_code == 409
        assert "Conflicting role reuse is not permitted" in conflicting.text


def test_original_source_file_on_disk_is_unchanged(client, auth_headers, monkeypatch):
    """Ensure imported source bytes are copies and original file on disk is not mutated."""
    client.put("/api/v1/career-profile", json=_profile_payload(0), headers=auth_headers)
    with TemporaryDirectory() as temp_dir:
        monkeypatch.setattr("backend.career.sources.settings.DATA_DIR", temp_dir)
        source_path = Path(temp_dir) / "OriginalFile.md"
        original_bytes = b"# Original File\n\nMust not be modified on disk."
        source_path.write_bytes(original_bytes)

        res = client.post(
            "/api/v1/career-profile/sources",
            headers=auth_headers,
            data={"source_role": "profile"},
            files={"file": ("OriginalFile.md", source_path.read_bytes(), "text/markdown")},
        )
        assert res.status_code == 201
        # Check that original file on disk remains completely identical
        assert source_path.read_bytes() == original_bytes


def test_unsafe_files_and_secrets_are_rejected():
    """Verify scripts, credentials, secret tokens, and control characters are rejected without logging secrets."""
    # 1. Script file extensions
    with pytest.raises(SourceImportError, match="Unsupported or unsafe file"):
        validate_safe_source_input("script.py", b"print('hello')")

    with pytest.raises(SourceImportError, match="Unsupported or unsafe file"):
        validate_safe_source_input("deploy.sh", b"echo 'run'")

    with pytest.raises(SourceImportError, match="Unsupported or unsafe file"):
        validate_safe_source_input("file.py.txt", b"x = 1")

    # 2. Credential filenames
    with pytest.raises(SourceImportError, match="Unsupported or credential file"):
        validate_safe_source_input("id_rsa", b"key data")

    with pytest.raises(SourceImportError, match="Unsupported or credential file"):
        validate_safe_source_input(".env", b"SECRET=foo")

    with pytest.raises(SourceImportError, match="Unsupported or credential file"):
        validate_safe_source_input("credentials.json", b"{}")

    # 3. Shebang scripts
    with pytest.raises(SourceImportError, match="Executable script files cannot be imported"):
        validate_safe_source_input("note.txt", b"#!/bin/bash\necho hi")

    # 4. Secret patterns in content
    with pytest.raises(SourceImportError, match="credential or private secret material"):
        validate_safe_source_input("key.txt", b"-----BEGIN RSA PRIVATE KEY-----\nMIIEowIBAAKCAQEA0...")

    with pytest.raises(SourceImportError, match="credential or private secret material"):
        validate_safe_source_input("token.txt", b"my token is ghp_123456789012345678901234567890123456")

    # 5. Control characters in text
    with pytest.raises(SourceImportError, match="unsupported binary or control characters"):
        validate_safe_source_input("binary.txt", b"Hello\x00World")

    with pytest.raises(SourceImportError, match="unsupported binary or control characters"):
        validate_safe_source_input("ctrl.md", b"Test\x07bell")


def test_invalid_preference_values_and_combined_conflicts():
    """Verify that invalid values are omitted with review notes and combined contradictions produce warnings."""
    # Invalid notice period > 730
    text_invalid_val = "Notice period days: 9999\nWorkload min: 150\nPreferred work modes: teleportation"
    candidates, notes, warnings = parse_goal_preferences(text_invalid_val)
    assert len(candidates) == 0
    assert any("notice_period_days" in note for note in notes)
    assert any("workload_min" in note for note in notes)
    assert any("teleportation" in note for note in notes)

    # Conflicting combined preferences: workload_min > workload_max
    text_conflict = "Workload min: 90\nWorkload max: 60"
    candidates, notes, warnings = parse_goal_preferences(text_conflict)
    assert len(candidates) == 2
    assert len(warnings) > 0
    assert any("workload_min cannot exceed workload_max" in w for w in warnings)

    # Remote only with onsite work mode conflict
    text_conflict_remote = "Remote only: true\nPreferred work modes: onsite"
    candidates, notes, warnings = parse_goal_preferences(text_conflict_remote)
    assert len(warnings) > 0
    assert any("remote_only conflicts with preferred_work_modes" in w for w in warnings)


def test_preference_candidate_contract_rejects_unknown_fields_and_invalid_values():
    base = {
        "candidate_id": "a" * 64,
        "source_locator": "line:1",
        "excerpt": "Synthetic preference",
    }
    with pytest.raises(ValidationError):
        PreferenceCandidate(field="job_source_consents", value={"job_room": True}, **base)
    with pytest.raises(ValidationError):
        PreferenceCandidate(field="workload_min", value=-1, **base)


def test_nested_and_merged_docx_tables_traversed_once_in_document_order():
    """Ensure DOCX text extraction includes nested/merged tables once in document order with bounds."""
    doc = Document()
    doc.add_paragraph("Top Paragraph 1")

    # Table with a horizontally merged cell
    table1 = doc.add_table(rows=2, cols=2)
    table1.cell(0, 0).merge(table1.cell(0, 1))
    table1.cell(0, 0).text = "Merged Cell Row 0"
    table1.cell(1, 0).text = "Cell R1C0"

    # Nested table inside cell R1C1
    nested_cell = table1.cell(1, 1)
    nested_cell.text = "Outer Cell Text"
    nested_table = nested_cell.add_table(rows=1, cols=1)
    nested_table.cell(0, 0).text = "Nested Table Evidence"

    doc.add_paragraph("Top Paragraph 2")

    bio = BytesIO()
    doc.save(bio)
    docx_bytes = bio.getvalue()

    extracted = extract_docx_text_in_document_order(docx_bytes, max_chars=10_000)

    # All text must be present
    assert "Top Paragraph 1" in extracted
    assert "Merged Cell Row 0" in extracted
    assert "Cell R1C0" in extracted
    assert "Outer Cell Text" in extracted
    assert "Nested Table Evidence" in extracted
    assert "Top Paragraph 2" in extracted

    # Merged cell text must appear exactly once
    assert extracted.count("Merged Cell Row 0") == 1

    # Order must be preserved
    pos_top1 = extracted.index("Top Paragraph 1")
    pos_merged = extracted.index("Merged Cell Row 0")
    pos_nested = extracted.index("Nested Table Evidence")
    pos_top2 = extracted.index("Top Paragraph 2")
    assert pos_top1 < pos_merged < pos_nested < pos_top2


def test_goal_import_bounds_review_notes_before_response_validation(client, auth_headers, monkeypatch):
    client.put("/api/v1/career-profile", json=_profile_payload(0), headers=auth_headers)
    content = "\n".join(f"Unrecognized goal prose number {index}" for index in range(51)).encode()
    with TemporaryDirectory() as temp_dir:
        monkeypatch.setattr("backend.career.sources.settings.DATA_DIR", temp_dir)
        response = client.post(
            "/api/v1/career-profile/sources",
            headers=auth_headers,
            data={"source_role": "goals"},
            files={"file": ("Goal.md", content, "text/markdown")},
        )

    assert response.status_code == 201, response.text
    notes = response.json()["review_notes"]
    assert len(notes) == 50
    assert "additional review notes omitted" in notes[-1]


@pytest.mark.parametrize(
    ("line", "field"),
    [
        ("Workload min: -20", "workload_min"),
        ("Workload min: 25.5", "workload_min"),
        ("Notice period days: -5", "notice_period_days"),
        ("Notice period days: later", "notice_period_days"),
        ("Max distance: -10", "hard_max_distance_km"),
        ("Max distance: nearby", "hard_max_distance_km"),
    ],
)
def test_goal_numeric_values_are_parsed_as_whole_values(line, field):
    candidates, notes, _warnings = parse_goal_preferences(line)

    assert candidates == []
    assert any(field in note and "not imported" in note for note in notes)


def test_goal_review_notes_are_bounded_with_an_explicit_omission_summary():
    text = "\n".join(f"Unrecognized narrative preference line number {index}" for index in range(51))

    _candidates, notes, _warnings = parse_goal_preferences(text)

    assert len(notes) == 50
    assert notes[-1] == "2 additional review notes omitted because the response limit was reached."


def test_conflicting_scalar_candidates_require_an_explicit_choice():
    candidates, _notes, warnings = parse_goal_preferences(
        "Workload min: 80\nWorkload min: 20"
    )

    assert [candidate.value for candidate in candidates] == [80, 20]
    assert any("choose exactly one" in warning for warning in warnings)


def test_large_docx_table_does_not_drop_cells_when_element_proxies_are_reused():
    doc = Document()
    table = doc.add_table(rows=40, cols=2)
    expected = []
    for row_index, row in enumerate(table.rows):
        for cell_index, cell in enumerate(row.cells):
            marker = f"marker-{row_index:02d}-{cell_index}"
            cell.text = marker
            expected.append(marker)
    output = BytesIO()
    doc.save(output)

    extracted = extract_docx_text_in_document_order(output.getvalue(), max_chars=20_000)

    assert all(extracted.count(marker) == 1 for marker in expected)
