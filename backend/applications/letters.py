"""Local paired A4 letters with escaped text and verified visible content."""

from __future__ import annotations

import io
import re
from datetime import datetime, timezone
from typing import Any, TypedDict
from urllib.parse import urlsplit
from xml.sax.saxutils import escape, quoteattr

from docx import Document
from docx.shared import Mm, Pt, RGBColor
from pypdf import PdfReader
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from backend.applications.placeholders import contains_unresolved_placeholder
from backend.applications.schemas import LetterOptions
from backend.resumes.renderers.fonts import ensure_unicode_font_registered
from backend.resumes.renderers.links import add_docx_hyperlink, is_safe_url
from backend.resumes.templates import get_template_preset

_URL = re.compile(r"(?:https?://|mailto:)[^\s<>\"']+")


class PageBudgetExceededError(ValueError):
    """A letter cannot fit its declared page budget."""


def validate_letter_text(text: str | None, *, field_name: str = "Cover letter") -> None:
    if text and (
        contains_unresolved_placeholder(text)
        or any(ord(c) < 32 and c not in "\n\t" for c in text)
    ):
        raise ValueError(f"{field_name} contains unresolved placeholders or invalid controls")


def _links(text: str):
    for match in _URL.finditer(text):
        url = match.group().rstrip(".,;)")
        try:
            parsed = urlsplit(url)
            safe = (
                is_safe_url(url)
                and not parsed.username
                and not parsed.password
                and (parsed.scheme == "mailto" or bool(parsed.hostname))
            )
        except ValueError:
            safe = False
        if safe:
            yield match.start(), match.start() + len(url), url


def _format_links_for_reportlab(text: str) -> str:
    pieces, previous = [], 0
    for start, end, url in _links(text):
        pieces.append(escape(text[previous:start]))
        pieces.append(f'<link href={quoteattr(url)} color="#1d4ed8">{escape(url)}</link>')
        previous = end
    pieces.append(escape(text[previous:]))
    return "".join(pieces).replace("\n", "<br/>")


def _preset(options: LetterOptions):
    preset = get_template_preset(options.preset_id, options.template_version)
    if options.locale != preset.locale:
        raise ValueError("Letter locale must match its selected template")
    return preset


def _parts(
    cover_letter: str,
    applicant_name: str,
    applicant_contact: list[str],
    recipient_lines: list[str],
    subject: str,
    date_str: str,
    locale: str,
):
    if not cover_letter.strip() or not applicant_name.strip():
        raise ValueError("Letter requires sender and body text")
    closing = "Mit freundlichen Grüssen," if locale == "de" else "Sincerely,"
    result = [
        (applicant_name, "name"),
        (" | ".join(filter(bool, applicant_contact)), "contact"),
        (date_str, "detail"),
        ("\n".join(filter(bool, recipient_lines)), "detail"),
        (subject, "subject"),
    ]
    result.extend((p.strip(), "body") for p in cover_letter.split("\n\n") if p.strip())
    result.extend([(closing, "closing"), (applicant_name, "signature")])
    for text, _ in result:
        validate_letter_text(text)
    return [(text, kind) for text, kind in result if text]


def render_letter_pdf(
    *,
    cover_letter: str,
    options: LetterOptions,
    applicant_name: str,
    applicant_contact: list[str],
    recipient_lines: list[str],
    subject: str,
    date_str: str,
    max_pages: int = 1,
) -> bytes:
    preset = _preset(options)
    regular, bold, _ = ensure_unicode_font_registered()
    margin = float(preset.preview_style.get("margin_mm", 20)) * 72 / 25.4
    parts = _parts(
        cover_letter,
        applicant_name,
        applicant_contact,
        recipient_lines,
        subject,
        date_str,
        options.locale,
    )
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=margin,
        bottomMargin=margin,
    )
    story: list[Any] = []
    for text, kind in parts:
        prominent = kind in {"name", "subject", "signature"}
        size = 15 if kind == "name" else 11 if kind == "subject" else 9.5
        style = ParagraphStyle(
            kind,
            fontName=bold if prominent else regular,
            fontSize=size,
            leading=size * 1.4,
            textColor=HexColor("#111827"),
            spaceAfter=8,
        )
        if kind == "closing":
            story.append(Spacer(1, 8))
        story.append(Paragraph(_format_links_for_reportlab(text), style))
    doc.build(story)
    data = buffer.getvalue()
    reader = PdfReader(io.BytesIO(data), strict=True)
    if not reader.pages or len(reader.pages) > max_pages:
        raise PageBudgetExceededError(f"Cover letter exceeds the {max_pages}-page A4 budget")
    for page in reader.pages:
        if (
            abs(float(page.mediabox.width) - A4[0]) > 1
            or abs(float(page.mediabox.height) - A4[1]) > 1
        ):
            raise ValueError("Letter PDF page is not A4")
    extracted = " ".join(" ".join(page.extract_text().split()) for page in reader.pages)
    if any(" ".join(text.split()) not in extracted for text, _ in parts):
        raise ValueError("Letter export did not preserve all visible text")
    return data


def render_letter_docx(
    *,
    cover_letter: str,
    options: LetterOptions,
    applicant_name: str,
    applicant_contact: list[str],
    recipient_lines: list[str],
    subject: str,
    date_str: str,
) -> bytes:
    preset = _preset(options)
    parts = _parts(
        cover_letter,
        applicant_name,
        applicant_contact,
        recipient_lines,
        subject,
        date_str,
        options.locale,
    )
    doc = Document()
    margin = float(preset.preview_style.get("margin_mm", 20))
    for section in doc.sections:
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.top_margin = section.bottom_margin = Mm(margin)
        section.left_margin = section.right_margin = Mm(margin)
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Arial", Pt(9.5)
    style.font.color.rgb = RGBColor.from_string("111827")
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.line_spacing = 1.4
    for text, kind in parts:
        paragraph = doc.add_paragraph()
        if kind == "closing":
            paragraph.paragraph_format.space_before = Pt(8)
        previous = 0
        for start, end, url in _links(text):
            paragraph.add_run(text[previous:start])
            add_docx_hyperlink(paragraph, url, url)
            previous = end
        paragraph.add_run(text[previous:])
        for run in paragraph.runs:
            run.bold = kind in {"name", "subject", "signature"}
            run.font.size = Pt(15 if kind == "name" else 11 if kind == "subject" else 9.5)
    buffer = io.BytesIO()
    doc.save(buffer)
    data = buffer.getvalue()
    reopened = Document(io.BytesIO(data))
    # Native XML text includes hyperlink runs, unlike paragraph.text on older python-docx.
    texts = [
        "".join(
            node.text or "" if node.tag.endswith("}t") else "\n"
            for node in paragraph._p.iter()
            if node.tag.endswith(("}t", "}br"))
        )
        for paragraph in reopened.paragraphs
    ]
    if any(
        " ".join(text.split()) not in " ".join(" ".join(v.split()) for v in texts)
        for text, _ in parts
    ):
        raise ValueError("Letter DOCX did not preserve all visible text")
    return data


class LetterRenderArguments(TypedDict):
    cover_letter: str
    options: LetterOptions
    applicant_name: str
    applicant_contact: list[str]
    recipient_lines: list[str]
    subject: str
    date_str: str


def generate_letter_artifacts(
    *,
    cover_letter: str | None,
    options: LetterOptions | None,
    applicant_name: str,
    applicant_contact: list[str],
    recipient_lines: list[str],
    subject: str,
    date_str: str | None = None,
) -> dict[str, tuple[bytes, str]]:
    if not options:
        validate_letter_text(cover_letter)
        return {}
    if not cover_letter or not cover_letter.strip():
        raise ValueError("Selected letter formats require cover letter text")
    values: LetterRenderArguments = dict(
        cover_letter=cover_letter.strip(),
        options=options,
        applicant_name=applicant_name,
        applicant_contact=applicant_contact,
        recipient_lines=recipient_lines,
        subject=options.subject or subject,
        date_str=date_str or options.date or datetime.now(timezone.utc).strftime("%Y-%m-%d"),
    )
    # Enforce the page budget even when only the editable format is selected.
    pdf = render_letter_pdf(**values)
    result = {}
    if "pdf" in options.formats:
        result["cover_letter.pdf"] = (pdf, "application/pdf")
    if "docx" in options.formats:
        result["cover_letter.docx"] = (
            render_letter_docx(**values),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    return result
