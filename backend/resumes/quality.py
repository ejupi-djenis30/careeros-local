from __future__ import annotations

import hashlib
import re
import zipfile
from io import BytesIO

from docx import Document
from docx.document import Document as DocxDocument
from docx.text.paragraph import Paragraph
from pypdf import PdfReader

from backend.core.config import settings
from backend.resumes.artifact_policy import (
    MAX_RESUME_ARTIFACT_BYTES,
    MAX_RESUME_DOCX_ENTRIES,
    MAX_RESUME_DOCX_UNCOMPRESSED_BYTES,
)

PLACEHOLDER_PATTERNS = [
    re.compile(
        r"\[(?:COMPANY|DATE|NAME|TITLE|ORGANIZATION|TODO|INSERT|ROLE|CITY|URL|EMAIL|PHONE|EMPLOYER|[A-Z0-9_ -]{2,})\]"
    ),
    re.compile(
        r"(?i)\[(?:company|date|name|title|organization|todo|insert|role|city|url|email|phone|employer)\]"
    ),
    re.compile(r"\b(?:TODO|FIXME):"),
]


class ResumeQualityError(ValueError):
    pass


def _normalized(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip().casefold()


def _in_order(text: str, values: list[str]) -> bool:
    cursor = 0
    for value in values:
        normalized = _normalized(value)
        if not normalized:
            continue
        position = text.find(normalized, cursor)
        if position < 0:
            return False
        cursor = position + len(normalized)
    return True


def check_no_placeholders(texts: list[str]) -> None:
    for text in texts:
        for pattern in PLACEHOLDER_PATTERNS:
            match = pattern.search(text)
            if match:
                raise ResumeQualityError(
                    f"Unresolved placeholder detected: '{match.group(0)}' in '{text[:80]}'"
                )


def extract_docx_text_in_order(document: DocxDocument) -> str:
    """Traverse all paragraphs and tables (including nested tables) in document order,
    skipping duplicate merged table cells."""
    extracted_lines: list[str] = []
    visited_cells: set = set()

    def _traverse(container, parent):
        for child in container:
            if child.tag.endswith("}p"):
                text = Paragraph(child, parent).text.strip()
                if text:
                    extracted_lines.append(text)
            elif child.tag.endswith("}tbl"):
                for tr in child:
                    if not tr.tag.endswith("}tr"):
                        continue
                    for tc in tr:
                        if not tc.tag.endswith("}tc"):
                            continue
                        v_merge = tc.xpath("./w:tcPr/w:vMerge")
                        if (
                            v_merge
                            and v_merge[0].get(
                                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val"
                            )
                            != "restart"
                        ):
                            continue
                        if tc in visited_cells:
                            continue
                        visited_cells.add(tc)
                        _traverse(tc, parent)

    _traverse(document.element.body, document)
    return "\n".join(extracted_lines)


def validate_resume_artifacts(
    *,
    pdf: bytes,
    docx: bytes,
    required_headings: list[str],
    required_text: list[str],
    template_kind: str,
    expect_photo: bool,
    columns: int = 1,
    page_budget: int | None = None,
) -> dict:
    # 1. Reject unresolved placeholders in requested content
    check_no_placeholders(required_headings + required_text)

    if len(pdf) > MAX_RESUME_ARTIFACT_BYTES or len(docx) > MAX_RESUME_ARTIFACT_BYTES:
        raise ResumeQualityError(
            f"Generated resume artifacts cannot exceed {MAX_RESUME_ARTIFACT_BYTES} bytes"
        )
    try:
        pdf_document = PdfReader(BytesIO(pdf))
        page_count = len(pdf_document.pages)
    except Exception as exc:
        raise ResumeQualityError("Generated PDF could not be reopened") from exc

    effective_max_pages = page_budget if page_budget else settings.RESUME_MAX_PAGES
    for page in pdf_document.pages:
        width, height = float(page.mediabox.width), float(page.mediabox.height)
        if abs(width - 595.276) > 2 or abs(height - 841.89) > 2:
            raise ResumeQualityError("Generated PDF must use portrait A4 on every page")
    if page_count < 1 or page_count > effective_max_pages:
        raise ResumeQualityError(
            f"Generated PDF has {page_count} pages; the configured limit is {effective_max_pages}"
        )

    try:
        extracted_text = "\n".join(page.extract_text() or "" for page in pdf_document.pages)
        pdf_image_count = sum(len(page.images) for page in pdf_document.pages)
        metadata = pdf_document.metadata
        pdf_metadata = {
            "author": metadata.author if metadata else None,
            "creator": metadata.creator if metadata else None,
            "producer": metadata.producer if metadata else None,
            "subject": metadata.subject if metadata else None,
        }
    except Exception as exc:
        raise ResumeQualityError("Generated PDF could not be reopened") from exc

    normalized_pdf = _normalized(extracted_text)
    missing_pdf = [
        item
        for item in required_headings + required_text
        if _normalized(item) not in normalized_pdf
    ]
    if missing_pdf:
        raise ResumeQualityError(
            "Generated PDF failed text extraction for: " + ", ".join(missing_pdf)
        )
    if not _in_order(normalized_pdf, required_headings) or not _in_order(
        normalized_pdf, required_text
    ):
        raise ResumeQualityError("Generated PDF does not preserve the requested text order")
    if template_kind == "ats" and pdf_image_count:
        raise ResumeQualityError("ATS PDF unexpectedly contains an image")
    if expect_photo and pdf_image_count < 1:
        raise ResumeQualityError("Photo PDF does not contain the normalized photo")

    try:
        with zipfile.ZipFile(BytesIO(docx)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_RESUME_DOCX_ENTRIES:
                raise ResumeQualityError("Generated DOCX contains too many archive entries")
            if len({entry.filename for entry in entries}) != len(entries):
                raise ResumeQualityError("Generated DOCX contains duplicate archive entries")
            if any(entry.flag_bits & 0x1 for entry in entries):
                raise ResumeQualityError("Generated DOCX cannot contain encrypted entries")
            if sum(entry.file_size for entry in entries) > MAX_RESUME_DOCX_UNCOMPRESSED_BYTES:
                raise ResumeQualityError("Generated DOCX expands beyond the local artifact limit")
            names = {entry.filename for entry in entries}
            if not {"[Content_Types].xml", "word/document.xml"} <= names:
                raise ResumeQualityError("Generated DOCX is missing required package entries")
            docx_image_count = sum(1 for name in names if name.startswith("word/media/"))
        word_document = Document(BytesIO(docx))

        # Check that EVERY section in DOCX explicitly uses A4 dimensions
        for sec_idx, sec in enumerate(word_document.sections):
            w_mm = sec.page_width.mm if sec.page_width else 0
            h_mm = sec.page_height.mm if sec.page_height else 0
            if abs(w_mm - 210.0) > 1.5 or abs(h_mm - 297.0) > 1.5:
                raise ResumeQualityError(
                    f"Generated DOCX section {sec_idx + 1} does not use A4 dimensions (found {w_mm:.1f}mm x {h_mm:.1f}mm)"
                )

        docx_text = extract_docx_text_in_order(word_document)
        docx_properties = word_document.core_properties
    except ResumeQualityError:
        raise
    except Exception as exc:
        raise ResumeQualityError("Generated DOCX could not be reopened") from exc

    normalized_docx = _normalized(docx_text)
    missing_docx = [
        item
        for item in required_headings + required_text
        if _normalized(item) not in normalized_docx
    ]
    if missing_docx:
        raise ResumeQualityError(
            "Generated DOCX is missing required content: " + ", ".join(missing_docx)
        )
    if not _in_order(normalized_docx, required_headings) or not _in_order(
        normalized_docx, required_text
    ):
        raise ResumeQualityError("Generated DOCX does not preserve the requested text order")
    if template_kind == "ats" and docx_image_count:
        raise ResumeQualityError("ATS DOCX unexpectedly contains an image")
    if expect_photo and docx_image_count < 1:
        raise ResumeQualityError("Photo DOCX does not contain the normalized photo")

    metadata_sanitized = (
        pdf_metadata.get("author") == "CareerOS Local"
        and pdf_metadata.get("creator") == "CareerOS Local"
        and pdf_metadata.get("producer") == "CareerOS Local"
        and pdf_metadata.get("subject") == "Resume"
        and docx_properties.author == "CareerOS Local"
        and docx_properties.comments == "Generated locally"
        and docx_properties.last_modified_by == "CareerOS Local"
    )
    if not metadata_sanitized:
        raise ResumeQualityError("Generated artifacts contain unexpected document metadata")

    return {
        "passed": True,
        "template_kind": template_kind,
        "layout": (
            "single-column"
            if template_kind == "ats"
            else f"{'two' if columns == 2 else 'single'}-column-photo"
        ),
        "page_count": page_count,
        "pdf_text_characters": len(extracted_text),
        "pdf_image_count": pdf_image_count,
        "docx_text_characters": len(docx_text),
        "docx_image_count": docx_image_count,
        "required_headings": required_headings,
        "text_order_verified": True,
        "metadata_sanitized": True,
        "within_page_limit": True,
        "max_pages": effective_max_pages,
        "pdf_sha256": hashlib.sha256(pdf).hexdigest(),
        "docx_sha256": hashlib.sha256(docx).hexdigest(),
    }
