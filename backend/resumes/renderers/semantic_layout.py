"""Swiss layouts with explicit reading order and bounded, local-only rendering."""

from io import BytesIO
from typing import Any

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.platypus import (
    BaseDocTemplate,
    Flowable,
    Frame,
    FrameBreak,
    Image,
    NextPageTemplate,
    PageBreak,
    PageTemplate,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from backend.resumes.artifact_policy import ensure_resume_artifact_size
from backend.resumes.content import build_content
from backend.resumes.quality import ResumeQualityError
from backend.resumes.renderers.base import (
    _add_docx_entry,
    _add_docx_heading,
    _canonicalize_docx,
    _configure_docx,
)
from backend.resumes.renderers.links import add_docx_hyperlink, format_pdf_link
from backend.resumes.renderers.photo_layout import _pdf_styles, _section_flowables, _text


def _header(content, styles, photo: bytes | None) -> list[Flowable]:
    flow: list[Flowable] = []
    if photo:
        picture = Image(BytesIO(photo), width=27 * mm, height=27 * mm)
        picture.hAlign = "LEFT"
        flow.extend([picture, Spacer(1, 8)])
    flow.append(Paragraph(_text(content.display_name), styles["name"]))
    if content.headline:
        flow.append(Paragraph(_text(content.headline), styles["headline"]))
    for contact in content.contact_line.split(" | "):
        if contact.strip():
            flow.append(Paragraph(format_pdf_link(contact.strip()), styles["contact"]))
    return flow


def _summary(content, styles) -> list[Flowable]:
    if not content.summary:
        return []
    return [
        Paragraph(_text(content.summary_heading), styles["section"]),
        Paragraph(_text(content.summary), styles["body"]),
    ]


def render_swiss_pdf(snapshot: dict, photo: bytes | None, *, operational: bool = False) -> bytes:
    content = build_content(snapshot)
    styles = _pdf_styles(content)
    for name in ("name", "headline", "contact"):
        styles[name].alignment = 0
    styles["name"].fontSize = 16 if not operational else 18
    styles["name"].leading = 20
    margin = float(content.style.get("margin_mm", 15)) * mm
    width, height = A4
    usable = width - 2 * margin
    output = BytesIO()
    options: dict[str, Any] = dict(
        pagesize=A4,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=margin,
        bottomMargin=margin,
        title=content.display_name,
        author="CareerOS Local",
        subject="Resume",
        creator="CareerOS Local",
        producer="CareerOS Local",
        invariant=True,
    )
    accent = HexColor(content.style.get("accent_color", "#243B53"))
    if operational:
        document: BaseDocTemplate = SimpleDocTemplate(output, **options)
        story = _header(content, styles, photo)
        story.extend(_summary(content, styles))
        for section in content.sections:
            if section.page_break_before:
                story.append(PageBreak())
            flow = _section_flowables(section, styles, group_entries=False)
            row = Table(
                [[flow[0], flow[1:]]],
                colWidths=[usable * 0.25, usable * 0.75],
                splitByRow=0,
                splitInRow=1,
                hAlign="LEFT",
            )
            row.setStyle(
                TableStyle(
                    [
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LINEBEFORE", (1, 0), (1, -1), 1, accent),
                        ("LEFTPADDING", (1, 0), (1, -1), 10),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                    ]
                )
            )
            story.append(row)
        document.build(story)
    else:
        sidebar_width = usable * 0.29
        gap = 9 * mm
        main_x = margin + sidebar_width + gap
        main_width = usable - sidebar_width - gap
        side = Frame(
            margin,
            margin,
            sidebar_width,
            height - 2 * margin,
            id="identity",
            leftPadding=0,
            rightPadding=0,
            topPadding=0,
            bottomPadding=0,
        )
        main = Frame(
            main_x,
            margin,
            main_width,
            height - 2 * margin,
            id="career",
            leftPadding=0,
            rightPadding=0,
            topPadding=0,
            bottomPadding=0,
        )

        def decorate(canvas, _document):
            canvas.saveState()
            canvas.setStrokeColor(accent)
            canvas.setLineWidth(1)
            canvas.line(main_x - gap / 2, margin, main_x - gap / 2, height - margin)
            canvas.restoreState()

        document = BaseDocTemplate(output, **options)
        document.addPageTemplates(
            [
                PageTemplate(id="first", frames=[side, main], onPage=decorate),
                PageTemplate(id="continuation", frames=[main], onPage=decorate),
            ]
        )
        header = _header(content, styles, photo)
        needed = sum(
            item.wrap(sidebar_width, height)[1] + item.getSpaceBefore() + item.getSpaceAfter()
            for item in header
        )
        if needed > height - 2 * margin:
            raise ResumeQualityError("Identity text exceeds the sidebar. Shorten it or choose ATS.")
        story = [
            NextPageTemplate("continuation"),
            *header,
            FrameBreak(),
            *_summary(content, styles),
        ]
        for section in content.sections:
            if section.page_break_before:
                story.append(PageBreak())
            story.extend(_section_flowables(section, styles))
        document.build(story)
    return ensure_resume_artifact_size(output.getvalue(), label="PDF")


def _docx_header(container, content, photo):
    if photo:
        container.add_paragraph().add_run().add_picture(BytesIO(photo), width=Inches(1.05))
    name = container.add_paragraph().add_run(content.display_name)
    name.bold = True
    name.font.size = Pt(18)
    if content.headline:
        container.add_paragraph(content.headline)
    for contact in content.contact_line.split(" | "):
        if contact.strip():
            add_docx_hyperlink(container.add_paragraph(), contact.strip())


def _shade(cell, color: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shading)


def render_swiss_docx(snapshot: dict, photo: bytes | None, *, operational: bool = False) -> bytes:
    content = build_content(snapshot)
    document = Document()
    _configure_docx(document, content.display_name, content.style)
    width_mm = 210 - 2 * float(content.style.get("margin_mm", 15))
    if operational:
        _docx_header(document, content, photo)
        main = document
    else:
        table = document.add_table(rows=1, cols=2)
        table.autofit = False
        table.columns[0].width = Mm(width_mm * 0.29)
        table.columns[1].width = Mm(width_mm * 0.71)
        side, main = table.rows[0].cells
        side.width, main.width = Mm(width_mm * 0.29), Mm(width_mm * 0.71)
        _shade(side, "F1F5F9")
        _docx_header(side, content, photo)
    if content.summary:
        _add_docx_heading(main, content.summary_heading)
        main.add_paragraph(content.summary)
    for section in content.sections:
        if section.page_break_before:
            main.add_paragraph().paragraph_format.page_break_before = True
        if operational:
            row = main.add_table(rows=1, cols=2)
            row.autofit = False
            row.columns[0].width, row.columns[1].width = Mm(width_mm * 0.25), Mm(width_mm * 0.75)
            heading_cell, entries_cell = row.rows[0].cells
            _shade(heading_cell, "F1F5F9")
            _add_docx_heading(heading_cell, section.heading)
        else:
            _add_docx_heading(main, section.heading)
            entries_cell = main
        for entry in section.entries:
            _add_docx_entry(entries_cell, entry)
    output = BytesIO()
    document.save(output)
    return _canonicalize_docx(output.getvalue())
